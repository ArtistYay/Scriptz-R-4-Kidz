"""Combine all CSV files in a directory into a single Excel workbook,
with each CSV file becoming its own worksheet.

Usage:
    python3 combine_csv_files_to_excel.py <csv_directory> <output_xlsx_path>
"""

from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
MAX_SHEET_TITLE_LENGTH = 31


def sanitize_value(value: str) -> str:
    """Strip characters that are not permitted in XLSX worksheet XML."""
    return ILLEGAL_CHARACTERS_RE.sub("", value)


def make_safe_sheet_title(name: str, used_titles: set) -> str:
    """Excel sheet names must be <=31 chars and cannot contain \\/*?:[]"""
    title = INVALID_SHEET_CHARS.sub("_", name)[:MAX_SHEET_TITLE_LENGTH]
    base_title = title
    suffix_number = 1
    while title.lower() in used_titles:
        suffix = f"_{suffix_number}"
        title = base_title[: MAX_SHEET_TITLE_LENGTH - len(suffix)] + suffix
        suffix_number += 1
    used_titles.add(title.lower())
    return title


def autosize_columns(worksheet: Worksheet, max_width: int = 60) -> None:
    for column_cells in worksheet.columns:
        longest_value = 0
        column_letter = None
        for cell in column_cells:
            if column_letter is None:
                column_letter = get_column_letter(cell.column)
            if cell.value is not None:
                longest_value = max(longest_value, len(str(cell.value)))
        if column_letter:
            worksheet.column_dimensions[column_letter].width = min(
                max(longest_value + 2, 10), max_width
            )


def write_csv_rows(worksheet: Worksheet, csv_path: Path) -> int:
    """Copy every CSV row into the worksheet and return how many were written."""
    row_count = 0
    with csv_path.open(newline="", encoding="utf-8-sig", errors="replace") as csv_file:
        reader = csv.reader(csv_file)
        for row in reader:
            worksheet.append([sanitize_value(cell) for cell in row])
            row_count += 1
    return row_count


def style_header_row(worksheet: Worksheet) -> None:
    """Freeze, filter and bold the first row."""
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True)


def add_csv_as_sheet(workbook: Workbook, csv_path: Path, used_titles: set) -> None:
    title = make_safe_sheet_title(csv_path.stem, used_titles)
    worksheet = workbook.create_sheet(title=title)

    row_count = write_csv_rows(worksheet, csv_path)
    if row_count > 0:
        style_header_row(worksheet)

    autosize_columns(worksheet)


def find_csv_files(csv_dir: Path) -> list[Path] | None:
    """Return the sorted CSV files in csv_dir, or None after printing an error."""
    if not csv_dir.is_dir():
        print(f"Error: '{csv_dir}' is not a directory", file=sys.stderr)
        return None

    csv_files = sorted(csv_dir.glob("*.csv"))
    if not csv_files:
        print(f"Error: no CSV files found in '{csv_dir}'", file=sys.stderr)
        return None
    return csv_files


def build_workbook(csv_files: list[Path]) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)  # remove default blank sheet

    used_titles: set = set()
    for csv_path in csv_files:
        print(f"Adding sheet for {csv_path.name}")
        add_csv_as_sheet(workbook, csv_path, used_titles)
    return workbook


def main() -> int:
    if len(sys.argv) != 3:
        print(
            f"Usage: {sys.argv[0]} <csv_directory> <output_xlsx_path>", file=sys.stderr
        )
        return 1

    csv_dir = Path(sys.argv[1]).expanduser().resolve()
    output_path = Path(sys.argv[2]).expanduser().resolve()

    csv_files = find_csv_files(csv_dir)
    if csv_files is None:
        return 1

    workbook = build_workbook(csv_files)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(f"Saved workbook with {len(csv_files)} sheets to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
